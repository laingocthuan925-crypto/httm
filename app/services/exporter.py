import io
from app.services.citations import render_citations, format_reference


def assemble(draft: dict, index: dict) -> dict:
    """Ghép bản thảo: trả {title, abstract, sections:[(heading,text)], references:[str]}."""
    used_all: list[str] = []
    sections = []
    for key in draft["order"]:
        s = draft["sections"][key]
        txt, used = render_citations(s["text"], index)
        used_all += [u for u in used if u not in used_all]
        sections.append((s["heading"], txt))
    refs = sorted((format_reference(index[i]) for i in used_all if i in index), key=str.lower)
    return {"title": draft.get("title", "Untitled"), "abstract": draft.get("abstract", ""),
            "sections": sections, "references": refs}


def to_markdown(doc: dict) -> str:
    out = [f"# {doc['title']}", "", "## Abstract", "", doc["abstract"], ""]
    for h, t in doc["sections"]:
        out += [f"## {h}", "", t, ""]
    out += ["## References", ""] + [f"- {r}" for r in doc["references"]]
    return "\n".join(out)


def to_docx(doc: dict) -> bytes:
    from docx import Document
    d = Document()
    d.add_heading(doc["title"], 0)
    d.add_heading("Abstract", 1)
    d.add_paragraph(doc["abstract"])
    for h, t in doc["sections"]:
        d.add_heading(h, 1)
        for para in [p.strip() for p in t.split("\n\n") if p.strip()]:
            d.add_paragraph(para)
    d.add_heading("References", 1)
    for r in doc["references"]:
        d.add_paragraph(r)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
